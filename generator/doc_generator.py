"""
Camada 4 — Geração do documento de briefing em Word (.docx).

Produz um documento estruturado por empresa com seção macro/setorial.
"""

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from config.settings import (
    OUTPUT_DIR,
    DOCUMENT_TITLE,
    DOCUMENT_SUBTITLE,
    COVERAGE_UNIVERSE,
)
from analyzer.api_analyzer import AnalysisResult

logger = logging.getLogger(__name__)


def _add_header(doc: Document, date_str: str) -> None:
    """Adiciona o header do documento."""
    # Título
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(DOCUMENT_TITLE)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 51, 102)

    # Subtítulo
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(DOCUMENT_SUBTITLE)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(100, 100, 100)

    # Data
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run(date_str)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    # Linha separadora
    doc.add_paragraph("─" * 80)


def _add_company_section(
    doc: Document,
    ticker: str,
    analyses: list[AnalysisResult],
) -> None:
    """Adiciona seção de uma empresa ao documento."""
    company_info = COVERAGE_UNIVERSE.get(ticker, {})
    company_name = company_info.get("company", ticker)
    sector = company_info.get("sector", "")

    # Header da empresa
    heading = doc.add_heading(level=2)
    run = heading.add_run(f"{ticker} — {company_name}")
    run.font.color.rgb = RGBColor(0, 51, 102)

    if sector:
        sector_para = doc.add_paragraph()
        run = sector_para.add_run(f"Setor: {sector}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

    for analysis in analyses:
        _add_news_item(doc, analysis)

    doc.add_paragraph()  # espaço entre seções


def _sentiment_indicator(sentiment: str) -> str:
    """Retorna indicador visual para o sentimento."""
    mapping = {
        "Positivo": "[+]",
        "Negativo": "[-]",
        "Neutro": "[=]",
    }
    return mapping.get(sentiment, "[=]")


def _add_news_item(doc: Document, analysis: AnalysisResult) -> None:
    """Adiciona um item de notícia ao documento."""
    # Título da notícia com sentimento
    indicator = _sentiment_indicator(analysis.sentiment)
    title_para = doc.add_paragraph()
    run = title_para.add_run(f"{indicator} {analysis.title}")
    run.bold = True
    run.font.size = Pt(10)

    # Sentimento colorido
    if analysis.sentiment == "Positivo":
        run.font.color.rgb = RGBColor(0, 128, 0)
    elif analysis.sentiment == "Negativo":
        run.font.color.rgb = RGBColor(192, 0, 0)
    else:
        run.font.color.rgb = RGBColor(64, 64, 64)

    # Fonte e data
    meta_para = doc.add_paragraph()
    run = meta_para.add_run(f"Fonte: {analysis.source} | {analysis.published_at[:10]}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(150, 150, 150)

    # Síntese
    summary_para = doc.add_paragraph()
    run = summary_para.add_run(analysis.summary)
    run.font.size = Pt(10)

    # Flag de ação requerida
    if analysis.action_required:
        action_para = doc.add_paragraph()
        run = action_para.add_run(f">> ACAO REQUERIDA: {analysis.action_detail or 'Revisão necessária'}")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(192, 0, 0)

    # Link
    link_para = doc.add_paragraph()
    run = link_para.add_run(analysis.url)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 102, 204)
    run.underline = True


def _add_summary_table(doc: Document, analyses: list[AnalysisResult]) -> None:
    """Adiciona tabela resumo no início do documento."""
    doc.add_heading("Resumo Executivo", level=1)

    # Conta por sentimento
    positive = sum(1 for a in analyses if a.sentiment == "Positivo")
    negative = sum(1 for a in analyses if a.sentiment == "Negativo")
    neutral = sum(1 for a in analyses if a.sentiment == "Neutro")
    actions = sum(1 for a in analyses if a.action_required)

    stats_para = doc.add_paragraph()
    stats_para.add_run(
        f"Total: {len(analyses)} notícias relevantes | "
        f"[+] {positive}  [-] {negative}  [=] {neutral}  "
        f"| Ações requeridas: {actions}"
    ).font.size = Pt(10)

    # Tabela de ações requeridas
    action_items = [a for a in analyses if a.action_required]
    if action_items:
        doc.add_paragraph()
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = table.rows[0].cells
        headers[0].text = "Ticker"
        headers[1].text = "Notícia"
        headers[2].text = "Sentimento"
        headers[3].text = "Ação"

        for item in action_items:
            row = table.add_row().cells
            row[0].text = ", ".join(item.tickers) if item.tickers else "Macro"
            row[1].text = item.title[:60]
            row[2].text = item.sentiment
            row[3].text = item.action_detail or "Revisão"

    doc.add_paragraph("─" * 80)


def generate_briefing(analyses: list[AnalysisResult]) -> Path:
    """
    Gera o documento Word de briefing.
    Retorna o path do arquivo gerado.
    """
    doc = Document()

    # Configura margens
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    date_str = datetime.now().strftime("%d/%m/%Y")

    _add_header(doc, date_str)

    if not analyses:
        doc.add_paragraph("Nenhuma notícia relevante identificada nas últimas 24 horas.")
        filename = f"briefing_{datetime.now().strftime('%Y%m%d')}.docx"
        output_path = OUTPUT_DIR / filename
        doc.save(str(output_path))
        logger.info(f"Documento gerado (vazio): {output_path}")
        return output_path

    # Tabela resumo
    _add_summary_table(doc, analyses)

    # Separa por empresa e macro
    company_analyses: dict[str, list[AnalysisResult]] = {}
    macro_analyses: list[AnalysisResult] = []

    for analysis in analyses:
        if analysis.category == "company" and analysis.tickers:
            for ticker in analysis.tickers:
                if ticker not in company_analyses:
                    company_analyses[ticker] = []
                company_analyses[ticker].append(analysis)
        else:
            macro_analyses.append(analysis)

    # Seções por empresa (ordenadas por ticker)
    if company_analyses:
        doc.add_heading("Notícias por Empresa", level=1)
        for ticker in sorted(company_analyses.keys()):
            _add_company_section(doc, ticker, company_analyses[ticker])

    # Seção macro/setorial
    if macro_analyses:
        doc.add_heading("Macro / Setorial", level=1)
        sector_para = doc.add_paragraph()
        run = sector_para.add_run(
            "Notícias relevantes para os setores de cobertura sem ticker específico."
        )
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

        for analysis in macro_analyses:
            _add_news_item(doc, analysis)

    # Footer
    doc.add_paragraph("─" * 80)
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        "Documento gerado automaticamente — Equity Research Pipeline\n"
        "As informações aqui contidas não constituem recomendação de investimento."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(150, 150, 150)

    # Salva
    filename = f"briefing_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
    output_path = OUTPUT_DIR / filename
    doc.save(str(output_path))
    logger.info(f"Documento gerado: {output_path}")
    return output_path
