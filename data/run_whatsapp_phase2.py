"""
Phase 2 — Real writes: WhatsApp clippings -> 3 production databases.

User confirmed: "Pode adicionar na base de dados do drive" (proceed with real writes).

Targets:
  DB1 — SQLite  output/news_database.db              key=md5(url)
  DB2 — JSONL   news_dataset/clipping_enriched.jsonl  key=sha256(url+title)[:32]
  DB3 — Excel   G:/My Drive/.../News_Clipping_Archive.xlsx  dedup=(date, headline)
"""

from __future__ import annotations

import hashlib
import json
import logging
import shutil
import sqlite3
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import requests
import torch

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

ROOT = Path(__file__).parent.parent
CHAT_PATH = Path(r"C:\Users\lucas\AppData\Local\Temp\WhatsApp Chat with Luiza Mussi.txt")
DB_PATH    = ROOT / "output" / "news_database.db"
JSONL_PATH = ROOT / "news_dataset" / "clipping_enriched.jsonl"
XLSX_PATH  = Path(r"G:\My Drive\4. Trabalho\Safra_CG_db\data\news\takes\News_Clipping_Archive.xlsx")
XLSX_SHEET = "Takes (proposed)"

sys.path.insert(0, str(ROOT))  # ensure project root is importable

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)-8s %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Section canonicalization (4 buckets matching BRIEFING_V2_SECTORS)
# ---------------------------------------------------------------------------

_SECTION_MAP: dict[str, str] = {
    # Aerospace & Defense
    "aerospace & defense":    "Aerospace & Defense",
    "aerospace and defense":  "Aerospace & Defense",
    "aerospace & defence":    "Aerospace & Defense",
    "aerospace and defence":  "Aerospace & Defense",
    "aerospace":              "Aerospace & Defense",
    "aviation":               "Aerospace & Defense",
    # Automotive
    "automotive":             "Automotive",
    "automotive & industrials": "Automotive",
    "automotivo":             "Automotive",
    "auto parts":             "Automotive",
    # Energy and Technology
    "energy and technology":  "Energy and Technology",
    "energy & technology":    "Energy and Technology",
    "energy and tech":        "Energy and Technology",
}


def _canon_section(raw: str) -> str:
    """Canonicalize a raw section name to one of the 4 standard buckets."""
    key = raw.strip().lower().rstrip(":")
    return _SECTION_MAP.get(key, "Others")


# ---------------------------------------------------------------------------
# Key helpers
# ---------------------------------------------------------------------------

def _db1_key(url: str) -> str:
    """MD5(url) — matches data/news_db.py:_article_key."""
    return hashlib.md5(url.encode()).hexdigest()


def _db2_key(url: str, title: str) -> str:
    """SHA-256(url+title)[:32] — matches data/ingest_history.py:_record_id."""
    return hashlib.sha256((url + title).encode()).hexdigest()[:32]


def _ddmmyyyy_to_iso(ddmmyyyy: str) -> str:
    """'15/01/2025' -> '2025-01-15'."""
    try:
        d, m, y = ddmmyyyy.split("/")
        return f"{y}-{int(m):02d}-{int(d):02d}"
    except Exception:
        return ddmmyyyy


# ---------------------------------------------------------------------------
# Pre-flight: service checks
# ---------------------------------------------------------------------------

def _check_ollama() -> bool:
    try:
        r = requests.get("http://localhost:11434/api/version", timeout=5)
        r.raise_for_status()
        log.info(f"Ollama reachable: {r.json().get('version')}")
        return True
    except Exception as e:
        log.warning(f"Ollama unreachable: {e}")
        return False


def _check_qdrant() -> bool:
    try:
        r = requests.get("http://localhost:6333/healthz", timeout=5)
        r.raise_for_status()
        log.info("Qdrant reachable")
        return True
    except Exception as e:
        log.warning(f"Qdrant unreachable: {e}")
        return False


# ---------------------------------------------------------------------------
# Backup helpers
# ---------------------------------------------------------------------------

def _backup(src: Path) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = src.suffix
    dest = src.with_name(src.stem + f"_backup_{ts}" + suffix)
    shutil.copy2(src, dest)
    log.info(f"Backup: {dest}")
    return dest


# ---------------------------------------------------------------------------
# DB1 helpers
# ---------------------------------------------------------------------------

def _db1_existing_keys(db_path: Path) -> set[str]:
    con = sqlite3.connect(str(db_path))
    cur = con.cursor()
    cur.execute("SELECT key FROM articles")
    keys = {row[0] for row in cur.fetchall()}
    con.close()
    return keys


def _db1_insert_batch(rows: list[dict], db_path: Path) -> int:
    """Insert rows into DB1. Each row must have keys matching the INSERT below."""
    inserted = 0
    with sqlite3.connect(str(db_path)) as con:
        for row in rows:
            cur = con.execute(
                """
                INSERT OR IGNORE INTO articles
                    (key, headline, url, source, date, published_at,
                     relevance_score, tier, capital_goods, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    row["key"],
                    row["headline"],
                    row["url"],
                    row["source"],
                    row["date"],
                    row["published_at"],
                    row["relevance_score"],
                    row["tier"],
                    1,  # capital_goods = 1 for all WhatsApp analyst-approved items
                    row["created_at"],
                ),
            )
            inserted += cur.rowcount
        con.commit()
    return inserted


# ---------------------------------------------------------------------------
# DB2 helpers
# ---------------------------------------------------------------------------

def _db2_existing_keys(jsonl_path: Path) -> set[str]:
    keys: set[str] = set()
    if not jsonl_path.exists():
        return keys
    with open(jsonl_path, encoding="utf-8", errors="replace") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            keys.add(_db2_key(r.get("url", "") or "", r.get("title", "") or ""))
    return keys


def _db2_append_batch(records: list[dict], jsonl_path: Path) -> int:
    with open(jsonl_path, "a", encoding="utf-8") as f:
        for r in records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    return len(records)


# ---------------------------------------------------------------------------
# DB3 helpers
# ---------------------------------------------------------------------------

def _db3_existing_pairs(xlsx_path: Path, sheet: str) -> tuple[set[tuple[str, str]], int]:
    """Return (existing (date, headline) pairs, max id)."""
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return set(), 0
    header = list(rows[0])
    date_idx = header.index("date")
    headline_idx = header.index("headline")
    id_idx = header.index("id")
    pairs: set[tuple[str, str]] = set()
    max_id = 0
    for row in rows[1:]:
        d = row[date_idx]
        h = row[headline_idx]
        raw_id = row[id_idx]
        if d is not None and h is not None:
            pairs.add((str(d).strip(), str(h).strip()))
        if raw_id is not None:
            try:
                max_id = max(max_id, int(str(raw_id).strip()))
            except (ValueError, TypeError):
                pass
    return pairs, max_id


def _db3_append_rows(new_rows: list[dict], xlsx_path: Path, sheet: str) -> int:
    """Append rows to the xlsx sheet. new_rows must have all column keys."""
    wb = openpyxl.load_workbook(str(xlsx_path))
    ws = wb[sheet]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    for row_dict in new_rows:
        row_vals = [row_dict.get(col) for col in header]
        ws.append(row_vals)
    wb.save(str(xlsx_path))
    wb.close()
    return len(new_rows)


# ---------------------------------------------------------------------------
# Scoring (BGE-M3 + Qdrant)
# ---------------------------------------------------------------------------

def _compute_scores_batch(
    items: list[dict],  # each has "text" key for embedding
    qdrant_available: bool,
) -> list[tuple[float, str]]:
    """
    Returns list of (score, tier) for each item.
    Falls back to score=1.0, tier="high" if Qdrant unavailable.
    """
    if not qdrant_available or not items:
        return [(1.0, "high")] * len(items)

    from sentence_transformers import SentenceTransformer
    from screener.embedding_store import EmbeddingStore
    from screener.relevance_scorer import (
        _compute_score, _assign_tier, _load_calibration
    )
    from config.settings import settings

    device = "cuda" if torch.cuda.is_available() else "cpu"
    log.info(f"Loading embedding model {settings.embedding_model!r} on {device}...")
    model = SentenceTransformer(settings.embedding_model, device=device)

    store = EmbeddingStore()
    if store.count() == 0:
        log.warning("Qdrant is empty — using fallback score=1.0 for all items.")
        return [(1.0, "high")] * len(items)

    cal = _load_calibration() or {}
    discard_threshold = cal.get("threshold", 0.40)
    tier_medium       = cal.get("tier_medium", 0.55)
    tier_high         = cal.get("tier_high", 0.85)
    k                 = cal.get("k", settings.embedding_top_k)
    k_neg             = cal.get("k_neg", settings.embedding_top_k_neg)

    texts = [it["text"] for it in items]
    log.info(f"Encoding {len(texts)} texts in batch...")
    embeddings: np.ndarray = model.encode(
        texts,
        batch_size=getattr(settings, "embedding_batch_size", 64),
        show_progress_bar=True,
        normalize_embeddings=True,
    )

    log.info("Querying Qdrant kNN for each item...")
    results: list[tuple[float, str]] = []
    for i, emb in enumerate(embeddings):
        if i % 500 == 0:
            log.info(f"  Scoring item {i}/{len(embeddings)}...")
        nearest_pos = store.search_knn(emb, k=k, label_filter="positive")
        nearest_neg = store.search_knn(emb, k=k_neg, label_filter="negative")
        score = _compute_score(nearest_pos, nearest_neg)
        # WhatsApp items are analyst-approved — clamp score to at least threshold
        score = max(score, discard_threshold)
        tier = _assign_tier(score, discard_threshold, tier_medium, tier_high)
        results.append((score, tier))

    return results


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    log.info("=" * 60)
    log.info("Phase 2: WhatsApp Clippings -> Production DBs (REAL WRITES)")
    log.info("=" * 60)

    # ------------------------------------------------------------------ #
    # 0. Pre-flight checks
    # ------------------------------------------------------------------ #
    ollama_ok = _check_ollama()
    qdrant_ok = _check_qdrant()

    # ------------------------------------------------------------------ #
    # 1. Count before
    # ------------------------------------------------------------------ #
    con = sqlite3.connect(str(DB_PATH))
    cur = con.cursor()
    cur.execute("SELECT COUNT(*) FROM articles")
    db1_before = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM articles WHERE capital_goods=1")
    db1_cg_before = cur.fetchone()[0]
    con.close()

    db2_before = sum(1 for line in open(JSONL_PATH, encoding="utf-8", errors="replace")
                     if line.strip())

    wb_tmp = openpyxl.load_workbook(str(XLSX_PATH), read_only=True, data_only=True)
    db3_before = wb_tmp[XLSX_SHEET].max_row - 1  # exclude header
    wb_tmp.close()

    log.info(f"BEFORE — DB1: {db1_before} rows ({db1_cg_before} capital_goods=1) | "
             f"DB2: {db2_before} JSONL | DB3: {db3_before} take rows")

    # ------------------------------------------------------------------ #
    # 2. Backups
    # ------------------------------------------------------------------ #
    log.info("Creating backups...")
    _backup(DB_PATH)
    _backup(JSONL_PATH)
    _backup(XLSX_PATH)

    # ------------------------------------------------------------------ #
    # 3. Parse WhatsApp export
    # ------------------------------------------------------------------ #
    from data.ingest_whatsapp_clippings import (
        parse_whatsapp, extract_clippings, extract_weekly_reminders,
        flatten_news, flatten_takes, extract_take_metadata_llm, resolve_url,
    )

    log.info(f"Parsing {CHAT_PATH}...")
    messages = parse_whatsapp(CHAT_PATH)
    clippings = extract_clippings(messages)
    weeklies  = extract_weekly_reminders(messages)
    log.info(f"Found {len(clippings)} clippings + {len(weeklies)} weekly reminders")

    # ------------------------------------------------------------------ #
    # 4. Flatten news — skip RESEARCH section
    # ------------------------------------------------------------------ #
    def _is_research(name: str) -> bool:
        n = name.strip().lower().rstrip(":")
        return n == "research" or n.startswith("research ")

    # Monkey-patch: filter RESEARCH section from each clipping before flatten
    for clip in clippings + weeklies:
        clip.sections = [s for s in clip.sections if not _is_research(s.name)]

    news_items = flatten_news(clippings, weeklies)
    takes      = flatten_takes(clippings, weeklies)
    log.info(f"Flattened: {len(news_items)} news items, {len(takes)} takes")

    # ------------------------------------------------------------------ #
    # 5. Canonicalize sections
    # ------------------------------------------------------------------ #
    for item in news_items:
        item.section = _canon_section(item.section)
    for take in takes:
        take.section = _canon_section(take.section)

    # ------------------------------------------------------------------ #
    # 6. Dedup against DB1 and DB2
    # ------------------------------------------------------------------ #
    log.info("Loading existing DB1 keys (md5(url))...")
    existing_db1 = _db1_existing_keys(DB_PATH)

    log.info("Loading existing DB2 keys (sha256(url+title)[:32])...")
    existing_db2 = _db2_existing_keys(JSONL_PATH)

    new_for_db1: list = []
    new_for_db2: list = []
    seen_for_scoring: dict[str, int] = {}  # url -> index in news_items

    for item in news_items:
        k1 = _db1_key(item.url)
        k2 = _db2_key(item.url, item.headline)
        if k1 not in existing_db1:
            new_for_db1.append(item)
        if k2 not in existing_db2:
            new_for_db2.append(item)

    log.info(f"New for DB1: {len(new_for_db1)} | New for DB2: {len(new_for_db2)}")

    # Union of items needing scoring (= new to DB1 — DB2 doesn't store score)
    items_to_score = new_for_db1
    score_texts = [
        f"{it.headline}\n\n{(it.body or '')[:4000]}" if it.body else it.headline
        for it in items_to_score
    ]
    score_inputs = [{"text": t} for t in score_texts]

    # ------------------------------------------------------------------ #
    # 7. Compute BGE-M3 + Qdrant scores
    # ------------------------------------------------------------------ #
    log.info(f"Scoring {len(items_to_score)} items...")
    score_results = _compute_scores_batch(score_inputs, qdrant_ok)

    # Build index: url -> (score, tier)  for quick lookup when writing
    url_to_score: dict[str, tuple[float, str]] = {}
    for item, (score, tier) in zip(items_to_score, score_results):
        url_to_score[item.url] = (score, tier)

    # ------------------------------------------------------------------ #
    # 8. Write DB1
    # ------------------------------------------------------------------ #
    log.info(f"Writing {len(new_for_db1)} rows to DB1...")
    now_iso = datetime.now(timezone.utc).isoformat()
    db1_rows = []
    for item in new_for_db1:
        score, tier = url_to_score.get(item.url, (1.0, "high"))
        date_iso = _ddmmyyyy_to_iso(item.date)
        db1_rows.append({
            "key":             _db1_key(item.url),
            "headline":        item.headline,
            "url":             item.url,
            "source":          "WhatsApp Clipping",
            "date":            date_iso,
            "published_at":    date_iso,
            "relevance_score": round(score, 6),
            "tier":            tier,
            "created_at":      now_iso,
        })

    db1_inserted = _db1_insert_batch(db1_rows, DB_PATH)
    log.info(f"DB1: {db1_inserted} rows inserted (of {len(new_for_db1)} candidates)")

    # ------------------------------------------------------------------ #
    # 9. Write DB2 JSONL
    # ------------------------------------------------------------------ #
    log.info(f"Writing {len(new_for_db2)} records to DB2 JSONL...")
    db2_records = []
    for item in new_for_db2:
        db2_records.append({
            "date":              item.date,
            "clipping_date_msg": item.date,   # msg time not preserved by flatten_news
            "sender":            item.sender,
            "section":           item.section,
            "title":             item.headline,
            "url":               item.url,
            "ticker":            None,
            "all_tickers":       [None],
            "relevant_yes_no":   True,
            "body":              item.body,
            "title_original":    None,
            "subtitle_original": None,
            "source_msg_type":   item.source_msg_type,
        })

    db2_appended = _db2_append_batch(db2_records, JSONL_PATH)
    log.info(f"DB2: {db2_appended} records appended")

    # ------------------------------------------------------------------ #
    # 10. Dedup takes against DB3
    # ------------------------------------------------------------------ #
    log.info("Loading existing DB3 take pairs...")
    existing_take_pairs, max_take_id = _db3_existing_pairs(XLSX_PATH, XLSX_SHEET)
    log.info(f"DB3: {len(existing_take_pairs)} existing pairs, max id = {max_take_id}")

    new_takes = []
    for take in takes:
        iso_date = _ddmmyyyy_to_iso(take.date)
        pair = (iso_date, take.headline)
        if pair not in existing_take_pairs:
            new_takes.append(take)
    log.info(f"New takes for DB3: {len(new_takes)}")

    # ------------------------------------------------------------------ #
    # 11. LLM extraction for new takes (parallel, 3 workers)
    # ------------------------------------------------------------------ #
    log.info(f"Running LLM extraction on {len(new_takes)} takes "
             f"({'Ollama available' if ollama_ok else 'FALLBACK: heuristic'})...")

    take_meta: list[dict] = [{}] * len(new_takes)

    def _extract_one(idx_take):
        idx, take = idx_take
        return idx, extract_take_metadata_llm(take)

    with ThreadPoolExecutor(max_workers=3) as pool:
        futures = {pool.submit(_extract_one, (i, t)): i for i, t in enumerate(new_takes)}
        done = 0
        for fut in as_completed(futures):
            idx, meta = fut.result()
            take_meta[idx] = meta
            done += 1
            if done % 50 == 0 or done == len(new_takes):
                log.info(f"  LLM extraction: {done}/{len(new_takes)} done")

    # ------------------------------------------------------------------ #
    # 12. Build DB3 rows (one row per company; fallback: one row with company="")
    # ------------------------------------------------------------------ #
    db3_rows: list[dict] = []
    next_id = max_take_id + 1

    for take, meta in zip(new_takes, take_meta):
        iso_date = _ddmmyyyy_to_iso(take.date)
        companies = meta.get("companies", [])
        topics    = meta.get("topics", "")

        if companies:
            for comp in companies:
                db3_rows.append({
                    "id":         str(next_id),
                    "date":       iso_date,
                    "company":    comp.get("name", ""),
                    "stance":     comp.get("stance", "Neutral"),
                    "topics":     topics,
                    "headline":   take.headline,
                    "take":       take.take_text,
                    "link":       take.url,
                    "marker":     take.marker,
                    "repeat_of":  None,
                    "flag":       None,
                    "approved":   "TRUE",
                })
                next_id += 1
        else:
            # No companies identified — still log the take
            db3_rows.append({
                "id":         str(next_id),
                "date":       iso_date,
                "company":    "",
                "stance":     "Neutral",
                "topics":     topics,
                "headline":   take.headline,
                "take":       take.take_text,
                "link":       take.url,
                "marker":     take.marker,
                "repeat_of":  None,
                "flag":       None,
                "approved":   "TRUE",
            })
            next_id += 1

    # ------------------------------------------------------------------ #
    # 13. Write DB3
    # ------------------------------------------------------------------ #
    log.info(f"Appending {len(db3_rows)} rows to DB3 xlsx...")
    db3_written = _db3_append_rows(db3_rows, XLSX_PATH, XLSX_SHEET)
    log.info(f"DB3: {db3_written} rows written")

    # ------------------------------------------------------------------ #
    # 14. Count after & summary
    # ------------------------------------------------------------------ #
    con = sqlite3.connect(str(DB_PATH))
    cur = con.cursor()
    cur.execute("SELECT COUNT(*) FROM articles")
    db1_after = cur.fetchone()[0]
    cur.execute("SELECT COUNT(*) FROM articles WHERE capital_goods=1")
    db1_cg_after = cur.fetchone()[0]
    con.close()

    db2_after = sum(1 for line in open(JSONL_PATH, encoding="utf-8", errors="replace")
                    if line.strip())

    wb_tmp = openpyxl.load_workbook(str(XLSX_PATH), read_only=True, data_only=True)
    db3_after = wb_tmp[XLSX_SHEET].max_row - 1
    wb_tmp.close()

    print("\n" + "=" * 60)
    print("PHASE 2 COMPLETE")
    print("=" * 60)
    print(f"DB1 (SQLite) : {db1_before} -> {db1_after} rows  "
          f"(+{db1_after - db1_before} new | capital_goods=1: {db1_cg_before} -> {db1_cg_after})")
    print(f"DB2 (JSONL)  : {db2_before} -> {db2_after} lines  "
          f"(+{db2_after - db2_before} new)")
    print(f"DB3 (Excel)  : {db3_before} -> {db3_after} rows  "
          f"(+{db3_after - db3_before} new | {len(new_takes)} unique takes -> {db3_written} xlsx rows)")
    print(f"Scoring      : {'BGE-M3 + Qdrant' if qdrant_ok else 'fallback score=1.0'}")
    print(f"LLM extract  : {'Ollama' if ollama_ok else 'heuristic fallback'}")
    print("=" * 60)


if __name__ == "__main__":
    main()
